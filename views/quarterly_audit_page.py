from datetime import date, datetime
from html import escape
from io import BytesIO
from urllib.parse import quote_plus

import streamlit as st

from database.db import (
    upsert_quarterly_audit_check,
)
from views.auth_page import get_current_user, require_user_id
from views.cached_data import (
    cached_get_quarterly_audit_checks,
    cached_get_quarterly_audit_history,
    cached_get_sites,
    clear_cached_data,
)


STATUS_OPTIONS = ["all_good", "needs_fix", "acceptable"]
STATUS_LABELS = {
    "all_good": "Все хорошо",
    "needs_fix": "Нужно исправить",
    "acceptable": "Допустимо",
}
STATUS_CSS = {
    "all_good": "qa-status-good",
    "needs_fix": "qa-status-fix",
    "acceptable": "qa-status-acceptable",
}
STATUS_COLORS = {
    "all_good": "C6EFCE",
    "needs_fix": "FCE4D6",
    "acceptable": "FFF2CC",
}

MAIN_CHECKS = [
    {
        "key": "responsive",
        "title": "Проверка адаптивности",
        "help": "Проверьте, удобно ли пользоваться основными страницами на мобильных устройствах.",
    },
    {
        "key": "site_speed",
        "title": "Скорость загрузки сайта",
        "help": "Зафиксируйте вручную Mobile score и Desktop score из PageSpeed.",
        "tools": [("PageSpeed", "https://pagespeed.web.dev/")],
        "speed": True,
    },
    {
        "key": "malware",
        "title": "Проверка вирусов",
        "help": "Проверьте сайт на вредоносный код, подозрительные скрипты и блокировки.",
        "tools": [("Dr.Web", "https://vms.drweb.ru/online")],
    },
    {
        "key": "cyclic_links",
        "title": "Циклические ссылки",
        "help": "Найдите страницы, где ссылки ведут сами на себя.",
        "tools": [("Screaming Frog", "https://www.screamingfrog.co.uk/seo-spider/")],
    },
    {
        "key": "extra_redirects",
        "title": "Лишние редиректы",
        "help": "Проверьте цепочки редиректов и оставьте только необходимые переходы.",
        "tools": [("Screaming Frog", "https://www.screamingfrog.co.uk/seo-spider/")],
    },
    {
        "key": "site_copies_index",
        "title": "Копии сайта в индексе",
        "help": "Проверьте дубли сайта, зеркала и копии контента в индексе.",
        "tools": [("Copyscape", "https://www.copyscape.com/")],
    },
    {
        "key": "commercial_factors",
        "title": "Проверка коммерческих факторов",
        "help": "Проверьте контакты, оплату, доставку, гарантии, отзывы и доверие.",
        "tools": [
            (
                "Открыть инструкцию",
                "https://sedate-frog-4ea.notion.site/57fee93802cb4098bed3e466b7bd6ae8?v=acc7b9598b924a32a8ccda536435d29e",
            )
        ],
    },
    {
        "key": "regionality",
        "title": "Проверка региональности в Вебмастере",
        "help": "Проверьте региональность сайта и локальные сигналы.",
        "tools": [("Яндекс Вебмастер", "https://webmaster.yandex.ru/")],
    },
    {
        "key": "filters",
        "title": "Проверка фильтров",
        "help": "Зафиксируйте выводы по фильтрам в комментарии.",
    },
    {
        "key": "image_alt",
        "title": "Alt картинок",
        "help": "Проверьте alt у важных изображений, карточек товаров и иллюстраций.",
    },
    {
        "key": "paid_links_indexation",
        "title": "Индексация закупленных ссылок",
        "help": "Проверьте индексацию страниц-доноров и статус размещений.",
    },
    {
        "key": "outgoing_links",
        "title": "Исходящие ссылки",
        "help": "Проверьте внешние ссылки, nofollow/sponsored и лишние сквозные ссылки.",
        "tools": [("Screaming Frog", "https://www.screamingfrog.co.uk/seo-spider/")],
    },
]

MONTHLY_CHECKS = [
    {
        "key": "webmaster_gsc_errors",
        "title": "Наличие ошибок в Яндекс Вебмастере и Google Search Console",
        "help": "Сверьте ошибки сканирования, индексации и покрытия в двух панелях.",
        "tools": [
            ("Яндекс Вебмастер", "https://webmaster.yandex.ru/"),
            ("Google Search Console", "https://search.google.com/search-console/"),
        ],
    },
    {"key": "iks", "title": "ИКС", "help": "Зафиксируйте ИКС и важные изменения в комментарии."},
    {"key": "page_indexation", "title": "Индексация страниц", "help": "Проверьте важные страницы в поиске."},
    {
        "key": "removed_pages_search",
        "title": "Удаленные страницы из поиска",
        "help": "Проверьте, какие страницы выпали из поиска и почему.",
    },
    {
        "key": "webmaster_summary",
        "title": "Сводка Вебмастера по сайту",
        "help": "Проверьте общую сводку сайта в Яндекс Вебмастере.",
        "tools": [("Яндекс Вебмастер", "https://webmaster.yandex.ru/")],
    },
    {
        "key": "sitemap_errors",
        "title": "Ошибки sitemap",
        "help": "Проверьте sitemap сайта и инструмент проверки sitemap в Вебмастере.",
        "tools": [
            ("Sitemap сайта", "{site_root}/sitemap.xml"),
            ("Проверка sitemap", "https://webmaster.yandex.ru/site/tools/sitemap/"),
        ],
    },
    {
        "key": "robots_txt",
        "title": "Robots",
        "help": "Проверьте robots.txt сайта и инструмент анализа robots в Вебмастере.",
        "tools": [
            ("Robots сайта", "{site_root}/robots.txt"),
            ("Проверка robots", "https://webmaster.yandex.ru/site/tools/robotstxt/"),
        ],
    },
    {
        "key": "duplicate_meta_pages",
        "title": "Дубли страниц и одинаковые мета-теги",
        "help": "Проверьте страницы с одинаковыми title/description и контентными дублями.",
    },
    {"key": "canonical", "title": "Canonical", "help": "Проверьте корректность canonical на типовых шаблонах."},
    {
        "key": "pagination",
        "title": "Пагинация",
        "help": "Проверка URL, title и meta-тегов на страницах пагинации. Нужно ли настроить rel canonical.",
    },
    {
        "key": "broken_links_404",
        "title": "Проверка битых ссылок 404",
        "help": "Зафиксируйте найденные 404 и план исправления.",
    },
    {
        "key": "important_meta_tags",
        "title": "Правильные мета-теги на важных страницах",
        "help": "Проверьте title и description на приоритетных посадочных страницах.",
    },
    {
        "key": "metrika_goals",
        "title": "Проверка целей в Метрике",
        "help": "Проверьте, что важные цели настроены и срабатывают.",
    },
    {
        "key": "homepage_duplicate_redirects",
        "title": "Главная страница должна редиректить",
        "help": "Проверьте index.php, /// и другие дубли главной страницы.",
        "tools": [("Проверка дублей", "https://be1.ru/dubli-stranic/")],
    },
    {
        "key": "mobile_core_pages",
        "title": "Проверка основных страниц, фильтра и корзины в мобильных устройствах",
        "help": "Проверьте сценарии на мобильных устройствах и зафиксируйте замечания.",
    },
    {
        "key": "reviews_links",
        "title": "Проверка отзывов",
        "help": "Используются ссылки на карты, сохраненные в разделе «Мои сайты».",
        "reviews": True,
    },
    {
        "key": "yandex_business_replies",
        "title": "Есть ли ответы компании на отзывы в Яндекс Бизнесе",
        "help": "Проверьте свежие отзывы и ответы компании.",
    },
    {
        "key": "yandex_business_photos",
        "title": "Разместить 2-3 фото в Яндекс Бизнесе",
        "help": "Зафиксируйте, какие фото добавлены или запланированы.",
    },
]

CHECK_SECTIONS = [
    ("Основной блок", MAIN_CHECKS),
    ("Ежемесячный аудит", MONTHLY_CHECKS),
]


def _site_name(site):
    return site[1] or site[2]


def _site_url(site):
    return site[2] or ""


def _site_root(site):
    raw_url = _site_url(site).strip()
    if not raw_url:
        return ""
    if raw_url.startswith(("http://", "https://")):
        return raw_url.rstrip("/")
    return f"https://{raw_url}".rstrip("/")


def _format_tool_url(url, site):
    root = _site_root(site)
    return url.format(
        site_url=quote_plus(root),
        site_root=root,
        site_host=quote_plus(root.replace("https://", "").replace("http://", "").split("/")[0]),
    )


def _site_review_links(site):
    candidates = [
        ("Яндекс Карты", site[5] if len(site) > 5 else ""),
        ("Google Maps", site[6] if len(site) > 6 else ""),
        ("2ГИС", site[7] if len(site) > 7 else ""),
    ]
    return [(label, url.strip()) for label, url in candidates if url and str(url).strip()]


def _tool_links(check, site):
    if check.get("reviews"):
        return _site_review_links(site)
    return [(label, _format_tool_url(url, site)) for label, url in check.get("tools", [])]


def _state_key(site_id, check_key, field):
    return f"quarterly_audit_{site_id}_{check_key}_{field}"


def _parse_date(value):
    if isinstance(value, date):
        return value
    if not value:
        return date.today()
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except ValueError:
        return date.today()


def _date_to_db(value):
    if isinstance(value, date):
        return value.isoformat()
    return str(value or date.today().isoformat())


def _score_color(value):
    if value is None:
        return "qa-score-empty"
    if value >= 90:
        return "qa-score-good"
    if value >= 50:
        return "qa-score-warn"
    return "qa-score-bad"


def _safe_int(value):
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _save_cell(user_id, site_id, check_key, user_name, has_speed=False):
    mobile_score = None
    desktop_score = None

    if has_speed:
        mobile_score = _safe_int(st.session_state.get(_state_key(site_id, check_key, "mobile_score")))
        desktop_score = _safe_int(st.session_state.get(_state_key(site_id, check_key, "desktop_score")))

    upsert_quarterly_audit_check(
        user_id=user_id,
        site_id=site_id,
        check_key=check_key,
        status=st.session_state.get(_state_key(site_id, check_key, "status"), "acceptable"),
        comment=st.session_state.get(_state_key(site_id, check_key, "comment"), ""),
        checked_at=_date_to_db(st.session_state.get(_state_key(site_id, check_key, "checked_at"))),
        checked_by=user_name,
        mobile_score=mobile_score,
        desktop_score=desktop_score,
        lcp=None,
        inp=None,
        cls=None,
    )
    clear_cached_data()
    st.session_state.pop("quarterly_export_data", None)
    st.session_state["quarterly_audit_saved_at"] = datetime.now().strftime("%H:%M:%S")


def _apply_date_to_site(user_id, site, selected_date, saved_checks, user_name):
    site_id = site[0]

    for check in _all_checks():
        check_key = check["key"]
        saved = saved_checks.get((site_id, check_key))
        _ensure_cell_state(site_id, check, saved)

        st.session_state[_state_key(site_id, check_key, "checked_at")] = selected_date
        upsert_quarterly_audit_check(
            user_id=user_id,
            site_id=site_id,
            check_key=check_key,
            status=st.session_state.get(_state_key(site_id, check_key, "status"), "acceptable"),
            comment=st.session_state.get(_state_key(site_id, check_key, "comment"), ""),
            checked_at=_date_to_db(selected_date),
            checked_by=user_name,
            mobile_score=_safe_int(st.session_state.get(_state_key(site_id, check_key, "mobile_score")))
            if check.get("speed")
            else None,
            desktop_score=_safe_int(st.session_state.get(_state_key(site_id, check_key, "desktop_score")))
            if check.get("speed")
            else None,
            lcp=None,
            inp=None,
            cls=None,
        )

    clear_cached_data()
    st.session_state.pop("quarterly_export_data", None)
    st.session_state["quarterly_audit_saved_at"] = datetime.now().strftime("%H:%M:%S")


def _ensure_cell_state(site_id, check, saved):
    check_key = check["key"]
    defaults = {
        "status": (saved.get("status") or "acceptable") if saved else "acceptable",
        "comment": saved.get("comment") if saved else "",
        "checked_at": _parse_date(saved.get("checked_at") if saved else None),
        "mobile_score": saved.get("mobile_score") if saved else None,
        "desktop_score": saved.get("desktop_score") if saved else None,
    }

    for field, value in defaults.items():
        key = _state_key(site_id, check_key, field)
        if key not in st.session_state:
            st.session_state[key] = value


def _status_badge(status):
    normalized = status if status in STATUS_OPTIONS else "acceptable"
    return f'<span class="qa-status {STATUS_CSS[normalized]}">{escape(STATUS_LABELS[normalized])}</span>'


def _all_checks():
    return [check for _, checks in CHECK_SECTIONS for check in checks]


def _check_title_map():
    return {check["key"]: check["title"] for check in _all_checks()}


def _status_totals(checks):
    totals = {"all_good": 0, "needs_fix": 0, "acceptable": 0}
    for check in checks.values():
        status = check.get("status") or "acceptable"
        if status in totals:
            totals[status] += 1
    return totals


def _render_css():
    st.markdown(
        """
        <style>
            .qa-hero {
                display: flex;
                justify-content: space-between;
                gap: 18px;
                align-items: flex-start;
                border: 1px solid #e2e8f0;
                border-radius: 10px;
                padding: 22px;
                margin-bottom: 16px;
                background: linear-gradient(135deg, rgba(14, 165, 233, 0.10), rgba(22, 163, 74, 0.08)), #ffffff;
                box-shadow: 0 10px 30px rgba(15, 23, 42, 0.06);
            }
            .qa-title {
                color: #0f172a;
                font-size: 30px;
                font-weight: 850;
                line-height: 1.15;
                margin-bottom: 8px;
            }
            .qa-subtitle {
                color: #475569;
                font-size: 14px;
                line-height: 1.55;
                max-width: 850px;
            }
            .qa-save-pill {
                display: inline-flex;
                white-space: nowrap;
                border-radius: 999px;
                padding: 7px 11px;
                color: #166534;
                background: #dcfce7;
                border: 1px solid #bbf7d0;
                font-size: 12px;
                font-weight: 850;
            }
            .qa-helper {
                color: #64748b;
                font-size: 13px;
                line-height: 1.45;
                margin: -2px 0 14px 0;
            }
            .qa-workspace-header,
            .qa-category {
                position: sticky;
                top: 0;
                z-index: 5;
                border: 1px solid #dbeafe;
                border-radius: 8px;
                padding: 10px;
                margin: 10px 0 12px 0;
                background: rgba(239, 246, 255, 0.98);
                backdrop-filter: blur(8px);
                box-shadow: 0 8px 18px rgba(15, 23, 42, 0.06);
            }
            .qa-category {
                display: flex;
                justify-content: space-between;
                align-items: center;
                background: rgba(248, 250, 252, 0.98);
                border-color: #e2e8f0;
                margin-top: 18px;
            }
            .qa-category-title {
                color: #0f172a;
                font-size: 18px;
                font-weight: 850;
            }
            .qa-category-note,
            .qa-workspace-head-cell {
                color: #1e3a8a;
                font-size: 12px;
                font-weight: 900;
                text-transform: uppercase;
                letter-spacing: 0.03em;
                overflow-wrap: anywhere;
            }
            .qa-row {
                border: 1px solid #e2e8f0;
                border-radius: 8px;
                background: #ffffff;
                margin-bottom: 10px;
                padding: 10px;
                box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04);
                transition: border-color 140ms ease, box-shadow 140ms ease;
            }
            .qa-row:hover {
                border-color: #bfdbfe;
                box-shadow: 0 10px 24px rgba(15, 23, 42, 0.07);
            }
            .qa-check-title {
                color: #0f172a;
                font-size: 15px;
                font-weight: 850;
                line-height: 1.3;
                margin-bottom: 7px;
            }
            .qa-check-help {
                color: #64748b;
                font-size: 12px;
                line-height: 1.42;
                margin-bottom: 10px;
            }
            .qa-site-chip {
                display: inline-flex;
                max-width: 100%;
                border-radius: 999px;
                padding: 4px 9px;
                color: #1d4ed8;
                background: #eff6ff;
                border: 1px solid #bfdbfe;
                font-size: 11px;
                font-weight: 850;
                overflow-wrap: anywhere;
            }
            .qa-status {
                display: inline-flex;
                align-items: center;
                border-radius: 999px;
                padding: 4px 9px;
                font-size: 11px;
                font-weight: 850;
                border: 1px solid transparent;
                margin-bottom: 6px;
            }
            .qa-status-good {
                color: #166534;
                background: #dcfce7;
                border-color: #bbf7d0;
            }
            .qa-status-fix {
                color: #9a3412;
                background: #ffedd5;
                border-color: #fed7aa;
            }
            .qa-status-acceptable {
                color: #854d0e;
                background: #fef9c3;
                border-color: #fde68a;
            }
            .qa-byline {
                color: #64748b;
                font-size: 11px;
                margin: 4px 0 6px 0;
            }
            .qa-speed-grid {
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 6px;
                margin: 4px 0 8px 0;
            }
            .qa-score {
                border-radius: 8px;
                padding: 6px 8px;
                font-size: 11px;
                font-weight: 850;
                border: 1px solid #e2e8f0;
            }
            .qa-score-good {
                color: #166534;
                background: #dcfce7;
                border-color: #bbf7d0;
            }
            .qa-score-warn {
                color: #9a3412;
                background: #ffedd5;
                border-color: #fed7aa;
            }
            .qa-score-bad {
                color: #991b1b;
                background: #fee2e2;
                border-color: #fecaca;
            }
            .qa-score-empty {
                color: #475569;
                background: #f8fafc;
                border-color: #e2e8f0;
            }
            .qa-empty {
                border: 1px dashed #cbd5e1;
                border-radius: 10px;
                padding: 18px;
                color: #64748b;
                background: #ffffff;
            }
            @media (max-width: 760px) {
                .qa-hero {
                    flex-direction: column;
                    padding: 18px;
                }
                .qa-title {
                    font-size: 24px;
                }
                .qa-workspace-header,
                .qa-category {
                    position: static;
                }
                .qa-category {
                    flex-direction: column;
                    align-items: flex-start;
                }
            }
        </style>
        """,
        unsafe_allow_html=True,
    )


def _render_speed_badges(site_id, check_key):
    mobile = _safe_int(st.session_state.get(_state_key(site_id, check_key, "mobile_score")))
    desktop = _safe_int(st.session_state.get(_state_key(site_id, check_key, "desktop_score")))

    st.markdown(
        f"""
        <div class="qa-speed-grid">
            <div class="qa-score {_score_color(mobile)}">Mobile: {mobile if mobile is not None else "-"}</div>
            <div class="qa-score {_score_color(desktop)}">Desktop: {desktop if desktop is not None else "-"}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_cell(user_id, site, check, saved, user_name):
    site_id = site[0]
    check_key = check["key"]
    has_speed = bool(check.get("speed"))
    _ensure_cell_state(site_id, check, saved)

    status = st.session_state.get(_state_key(site_id, check_key, "status"), "acceptable")
    checked_by = (saved or {}).get("checked_by") or user_name
    updated_at = (saved or {}).get("updated_at")
    callback_args = (user_id, site_id, check_key, user_name, has_speed)

    st.markdown(
        f"""
        <div>
            <div class="qa-site-chip">{escape(_site_name(site))}</div><br>
            {_status_badge(status)}
            <div class="qa-byline">SEO специалист: {escape(checked_by)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.selectbox(
        "Статус",
        STATUS_OPTIONS,
        format_func=lambda option: STATUS_LABELS[option],
        key=_state_key(site_id, check_key, "status"),
        label_visibility="collapsed",
        on_change=_save_cell,
        args=callback_args,
    )
    st.text_area(
        "Комментарий SEO специалиста",
        key=_state_key(site_id, check_key, "comment"),
        height=94,
        placeholder="Что проверено, что исправить, куда вернуться в следующем квартале...",
        label_visibility="collapsed",
        on_change=_save_cell,
        args=callback_args,
    )
    st.date_input(
        "Дата проверки",
        key=_state_key(site_id, check_key, "checked_at"),
        format="DD.MM.YYYY",
        on_change=_save_cell,
        args=callback_args,
    )

    if has_speed:
        _render_speed_badges(site_id, check_key)
        speed_cols = st.columns(2)
        with speed_cols[0]:
            st.number_input(
                "Mobile score",
                min_value=0,
                max_value=100,
                step=1,
                key=_state_key(site_id, check_key, "mobile_score"),
                on_change=_save_cell,
                args=callback_args,
            )
        with speed_cols[1]:
            st.number_input(
                "Desktop score",
                min_value=0,
                max_value=100,
                step=1,
                key=_state_key(site_id, check_key, "desktop_score"),
                on_change=_save_cell,
                args=callback_args,
            )

    if updated_at:
        st.caption(f"Сохранено: {updated_at}")

    links = _tool_links(check, site)
    if links:
        for index, (label, url) in enumerate(links):
            st.link_button(
                label,
                url,
                key=f"quarterly_tool_{site_id}_{check_key}_{index}",
                use_container_width=True,
            )
    else:
        st.caption("Без внешней ссылки. Заполните комментарий вручную.")


def _render_workspace_section(title, checks, visible_sites, sites_by_id, saved_checks, user_id, user_name):
    st.markdown(
        f"""
        <div class="qa-category">
            <div class="qa-category-title">{escape(title)}</div>
            <div class="qa-category-note">{len(checks)} проверок · {len(visible_sites)} сайтов</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    for check in checks:
        st.markdown('<div class="qa-row">', unsafe_allow_html=True)
        columns = st.columns([1.15] + [1 for _ in visible_sites], gap="medium")

        with columns[0]:
            st.markdown(
                f"""
                <div class="qa-check-title">{escape(check["title"])}</div>
                <div class="qa-check-help">{escape(check["help"])}</div>
                """,
                unsafe_allow_html=True,
            )
            st.caption("Ячейка сохраняется сразу после изменения поля.")

        for column, site in zip(columns[1:], visible_sites):
            with column:
                _render_cell(
                    user_id=user_id,
                    site=site,
                    check=check,
                    saved=saved_checks.get((site[0], check["key"])),
                    user_name=user_name,
                )

        st.markdown("</div>", unsafe_allow_html=True)


def _render_history(user_id, sites_by_id):
    history = cached_get_quarterly_audit_history(user_id=user_id, limit=60)
    if not history:
        st.markdown(
            '<div class="qa-empty">История появится после первых сохранений в таблице.</div>',
            unsafe_allow_html=True,
        )
        return

    check_titles = _check_title_map()
    for row in history[:20]:
        _, _, site_id, check_key, old_status, new_status, old_comment, new_comment, changed_at, changed_by = row
        site = sites_by_id.get(site_id)
        site_name = _site_name(site) if site else f"Сайт #{site_id}"
        title = check_titles.get(check_key, check_key)
        old_label = STATUS_LABELS.get(old_status, old_status or "пусто")
        new_label = STATUS_LABELS.get(new_status, new_status or "пусто")

        st.markdown(
            f"""
            <div class="ts-card" style="margin-bottom: 8px;">
                <div class="ts-card-title">{escape(site_name)} · {escape(title)}</div>
                <div class="ts-card-text">Статус: {escape(old_label)} → {escape(new_label)}</div>
                <div class="ts-card-text">Изменил: {escape(changed_by or "не указан")} · {escape(changed_at or "")}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        if (old_comment or "") != (new_comment or ""):
            st.caption("Комментарий был обновлен.")


def _export_quarterly_excel(sites, checks_data, history_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Проверки"
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False

    dark = "0F172A"
    blue = "2563EB"
    border = Side(style="thin", color="CBD5E1")
    header_fill = PatternFill("solid", fgColor=dark)
    section_fill = PatternFill("solid", fgColor="E0F2FE")
    muted_fill = PatternFill("solid", fgColor="F8FAFC")

    ws.merge_cells("A1:H1")
    ws["A1"] = "TechSEO Monitor"
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    ws["A1"].fill = PatternFill("solid", fgColor=blue)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "Quarterly SEO Workspace"
    ws["A2"].font = Font(color=dark, bold=True, size=14)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws["A3"] = "Дата экспорта"
    ws["B3"] = datetime.now().strftime("%d.%m.%Y %H:%M")
    ws["A4"] = "Сайты"
    ws["B4"] = ", ".join(_site_name(site) for site in sites)

    headers = [
        "Раздел",
        "Проверка",
        "Сайт",
        "Статус",
        "Комментарий",
        "Дата",
        "SEO специалист",
        "Mobile score",
        "Desktop score",
    ]
    ws.append(headers)
    for cell in ws[5]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 6
    for section, checks in CHECK_SECTIONS:
        ws.append([section])
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
        section_cell = ws.cell(row_num, 1)
        section_cell.fill = section_fill
        section_cell.font = Font(color=dark, bold=True)
        row_num += 1

        for check in checks:
            for site in sites:
                saved = checks_data.get((site[0], check["key"]), {})
                status = saved.get("status") or "acceptable"
                ws.append(
                    [
                        section,
                        check["title"],
                        _site_name(site),
                        STATUS_LABELS.get(status, status),
                        saved.get("comment") or "",
                        saved.get("checked_at") or "",
                        saved.get("checked_by") or "",
                        saved.get("mobile_score") if check.get("speed") else "",
                        saved.get("desktop_score") if check.get("speed") else "",
                    ]
                )
                status_cell = ws.cell(row_num, 4)
                status_cell.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(status, "FFFFFF"))
                row_num += 1

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=border, right=border, top=border, bottom=border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [22, 38, 24, 18, 48, 14, 22, 14, 14]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.auto_filter.ref = f"A5:I{ws.max_row}"

    summary = wb.create_sheet("Сводка")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "TechSEO Monitor"
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].fill = PatternFill("solid", fgColor=blue)
    summary["A2"] = "Сводка по статусам"
    summary["A2"].font = Font(color=dark, bold=True, size=14)
    summary.append(["Сайт", "Все хорошо", "Допустимо", "Нужно исправить", "Всего заполнено"])

    for site in sites:
        site_checks = [checks_data.get((site[0], check["key"])) for check in _all_checks()]
        site_checks = [item for item in site_checks if item]
        summary.append(
            [
                _site_name(site),
                sum(1 for item in site_checks if item.get("status") == "all_good"),
                sum(1 for item in site_checks if item.get("status") == "acceptable"),
                sum(1 for item in site_checks if item.get("status") == "needs_fix"),
                len(site_checks),
            ]
        )

    for cell in summary[3]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, 6):
        summary.column_dimensions[get_column_letter(col)].width = [28, 14, 14, 16, 16][col - 1]

    history = wb.create_sheet("История")
    history.sheet_view.showGridLines = False
    history.append(["Сайт", "Проверка", "Старый статус", "Новый статус", "Старый комментарий", "Новый комментарий", "Дата", "Кто изменил"])
    for cell in history[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    sites_by_id = {site[0]: site for site in sites}
    titles = _check_title_map()
    for row in history_rows:
        _, _, site_id, check_key, old_status, new_status, old_comment, new_comment, changed_at, changed_by = row
        site = sites_by_id.get(site_id)
        history.append(
            [
                _site_name(site) if site else f"Сайт #{site_id}",
                titles.get(check_key, check_key),
                STATUS_LABELS.get(old_status, old_status or ""),
                STATUS_LABELS.get(new_status, new_status or ""),
                old_comment or "",
                new_comment or "",
                changed_at or "",
                changed_by or "",
            ]
        )

    for row in history.iter_rows(min_row=1, max_row=history.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(left=border, right=border, top=border, bottom=border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, width in enumerate([24, 38, 18, 18, 42, 42, 20, 22], start=1):
        history.column_dimensions[get_column_letter(idx)].width = width

    for sheet in [ws, summary, history]:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row > 1 and cell.fill.fill_type is None:
                    cell.fill = muted_fill if cell.row % 2 == 0 else PatternFill(fill_type=None)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _export_quarterly_excel(sites, checks_data, history_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    status_labels = {
        "all_good": "🟢 Всё хорошо",
        "needs_fix": "🟠 Требует правок",
        "acceptable": "🟡 Пока терпит",
    }
    status_fills = {
        "all_good": "D9EAD3",
        "needs_fix": "FCE4D6",
        "acceptable": "FFF2CC",
    }
    status_font_colors = {
        "all_good": "166534",
        "needs_fix": "9A3412",
        "acceptable": "854D0E",
    }

    def parse_date(value):
        if isinstance(value, date):
            return value
        if not value:
            return None
        text = str(value).strip()
        for fmt, size in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10), ("%d.%m.%Y", 10)):
            try:
                return datetime.strptime(text[:size], fmt).date()
            except ValueError:
                continue
        return None

    def safe_sheet_name(name, used_names):
        invalid = set('[]:*?/\\')
        clean = "".join("_" if char in invalid else char for char in str(name or "Сайт")).strip()
        clean = clean.replace("https://", "").replace("http://", "").replace("/", "")
        clean = clean[:31] or "Сайт"
        candidate = clean
        counter = 2
        while candidate in used_names:
            suffix = f"_{counter}"
            candidate = f"{clean[:31 - len(suffix)]}{suffix}"
            counter += 1
        used_names.add(candidate)
        return candidate

    def tool_text_and_url(check, site):
        links = _tool_links(check, site)
        if not links:
            return "Без ссылки", None
        return "\n".join(label for label, _ in links), links[0][1]

    def make_cell_text(check, record, column_date):
        if not record:
            return ""

        comment = record.get("comment") or ""
        if check.get("speed"):
            mobile = record.get("mobile_score")
            desktop = record.get("desktop_score")
            return (
                f"Mobile: {mobile if mobile is not None else ''}\n"
                f"Desktop: {desktop if desktop is not None else ''}\n"
                f"Комментарий: {comment}"
            )

        status = record.get("status") or "acceptable"
        return (
            f"Статус: {status_labels.get(status, status)}\n"
            f"Комментарий: {comment}\n"
            f"Дата: {column_date.strftime('%d.%m.%Y')}"
        )

    def history_record_map(site_id):
        records = {}
        for row in history_rows:
            _, _, row_site_id, check_key, _, new_status, _, new_comment, changed_at, changed_by = row
            if row_site_id != site_id:
                continue
            changed_date = parse_date(changed_at)
            if not changed_date:
                continue
            records[(check_key, changed_date)] = {
                "status": new_status or "acceptable",
                "comment": new_comment or "",
                "checked_by": changed_by or "",
                "changed_at": changed_at or "",
            }
        return records

    wb = Workbook()
    wb.remove(wb.active)

    dark = "0F172A"
    blue = "2563EB"
    border = Side(style="thin", color="CBD5E1")
    header_fill = PatternFill("solid", fgColor=dark)
    section_fill = PatternFill("solid", fgColor="E0F2FE")
    muted_fill = PatternFill("solid", fgColor="F8FAFC")
    used_sheet_names = set()
    project_summaries = []
    titles = _check_title_map()

    for site in sites:
        site_id = site[0]
        site_name = _site_name(site)
        ws = wb.create_sheet(safe_sheet_name(site_name, used_sheet_names))
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C8"

        site_history = history_record_map(site_id)
        site_dates = set()
        for check in _all_checks():
            saved = checks_data.get((site_id, check["key"]))
            checked_date = parse_date(saved.get("checked_at") if saved else None)
            if checked_date:
                site_dates.add(checked_date)
        site_dates.update(date_value for _, date_value in site_history.keys())
        if not site_dates:
            site_dates.add(date.today())
        site_dates = sorted(site_dates)

        max_col = 2 + len(site_dates)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws["A1"] = "SEO-аудит 1 раз в месяц"
        ws["A1"].font = Font(color="FFFFFF", bold=True, size=18)
        ws["A1"].fill = PatternFill("solid", fgColor=blue)
        ws["A1"].alignment = Alignment(horizontal="center")

        ws["A2"] = "Проект"
        ws["B2"] = site_name
        ws["A3"] = "Дата формирования файла"
        ws["B3"] = datetime.now().strftime("%d.%m.%Y %H:%M")
        ws["A4"] = "Легенда"
        ws["B4"] = "🟢 Всё хорошо    🟠 Требует правок    🟡 Пока терпит"

        for row_index in range(2, 5):
            ws.cell(row_index, 1).font = Font(bold=True, color=dark)
            ws.cell(row_index, 1).fill = muted_fill
            ws.cell(row_index, 2).alignment = Alignment(vertical="center", wrap_text=True)

        header_row = 7
        ws.cell(header_row, 1, "Проверка")
        ws.cell(header_row, 2, "Инструмент")
        for index, check_date in enumerate(site_dates, start=3):
            ws.cell(header_row, index, check_date.strftime("%d.%m.%y"))

        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=border, right=border, top=border, bottom=border)

        row_num = header_row + 1
        for section, checks in CHECK_SECTIONS:
            ws.cell(row_num, 1, section)
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=max_col)
            section_cell = ws.cell(row_num, 1)
            section_cell.fill = section_fill
            section_cell.font = Font(color=dark, bold=True)
            section_cell.alignment = Alignment(vertical="center")
            row_num += 1

            for check in checks:
                ws.cell(row_num, 1, check["title"])
                tool_text, tool_url = tool_text_and_url(check, site)
                tool_cell = ws.cell(row_num, 2, tool_text)
                if tool_url:
                    tool_cell.hyperlink = tool_url
                    tool_cell.style = "Hyperlink"

                for date_index, check_date in enumerate(site_dates, start=3):
                    saved = checks_data.get((site_id, check["key"]))
                    record = None
                    if saved and parse_date(saved.get("checked_at")) == check_date:
                        record = dict(saved)
                    if (check["key"], check_date) in site_history:
                        record = {**(record or {}), **site_history[(check["key"], check_date)]}

                    cell = ws.cell(row_num, date_index, make_cell_text(check, record, check_date))
                    status = (record or {}).get("status")
                    if status in status_fills:
                        cell.fill = PatternFill("solid", fgColor=status_fills[status])
                        cell.font = Font(color=status_font_colors[status])
                    elif cell.value:
                        cell.fill = muted_fill

                row_num += 1

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = Border(left=border, right=border, top=border, bottom=border)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for row_index in range(8, ws.max_row + 1):
            ws.row_dimensions[row_index].height = 72
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[7].height = 30
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 24
        for col in range(3, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 28
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{ws.max_row}"

        latest_date = max(site_dates) if site_dates else None
        site_checks = [checks_data.get((site_id, check["key"])) for check in _all_checks()]
        site_checks = [item for item in site_checks if item]
        project_summaries.append(
            [
                site_name,
                latest_date.strftime("%d.%m.%Y") if latest_date else "",
                len(site_checks),
                sum(1 for item in site_checks if item.get("status") == "needs_fix"),
                sum(1 for item in site_checks if item.get("status") == "acceptable"),
            ]
        )

    summary = wb.create_sheet("Сводка", 0)
    summary.sheet_view.showGridLines = False
    summary["A1"] = "TechSEO Monitor"
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].fill = PatternFill("solid", fgColor=blue)
    summary["A2"] = "Сводка по проектам"
    summary["A2"].font = Font(color=dark, bold=True, size=14)
    summary.append(["Проект", "Последняя дата проверки", "Количество проверок", "Количество проблем", "Количество предупреждений"])
    for item in project_summaries:
        summary.append(item)

    for cell in summary[3]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in summary.iter_rows(min_row=1, max_row=summary.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = Border(left=border, right=border, top=border, bottom=border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate([32, 22, 22, 20, 24], start=1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = "A4"

    history = wb.create_sheet("История")
    history.sheet_view.showGridLines = False
    history.append(["Дата изменения", "Проект", "Проверка", "Старый статус", "Новый статус", "Старый комментарий", "Новый комментарий", "Кто изменил"])
    for cell in history[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    sites_by_id = {site[0]: site for site in sites}
    for row in history_rows:
        _, _, site_id, check_key, old_status, new_status, old_comment, new_comment, changed_at, changed_by = row
        site = sites_by_id.get(site_id)
        history.append(
            [
                changed_at or "",
                _site_name(site) if site else f"Сайт #{site_id}",
                titles.get(check_key, check_key),
                status_labels.get(old_status, old_status or ""),
                status_labels.get(new_status, new_status or ""),
                old_comment or "",
                new_comment or "",
                changed_by or "",
            ]
        )

    for row in history.iter_rows(min_row=1, max_row=history.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(left=border, right=border, top=border, bottom=border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, width in enumerate([22, 28, 42, 20, 20, 42, 42, 22], start=1):
        history.column_dimensions[get_column_letter(idx)].width = width
    history.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def show_quarterly_audit_page():
    user_id = require_user_id()
    user = get_current_user() or {}
    user_name = user.get("name") or user.get("email") or "SEO специалист"
    sites = cached_get_sites(user_id=user_id)
    sites_by_id = {site[0]: site for site in sites}
    saved_checks = cached_get_quarterly_audit_checks(user_id=user_id)

    _render_css()

    saved_at = st.session_state.get("quarterly_audit_saved_at")
    save_label = f"Автосохранено в {saved_at}" if saved_at else "Автосохранение включено"

    st.markdown(
        f"""
        <div class="qa-hero">
            <div>
                <div class="qa-title">Quarterly SEO Workspace</div>
                <div class="qa-subtitle">
                    Ручное рабочее пространство SEO-специалиста: основные квартальные проверки открыты сразу,
                    ежемесячный аудит скрыт в отдельном разделе, все изменения сохраняются в SQLite.
                </div>
            </div>
            <div class="qa-save-pill">{escape(save_label)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if not sites:
        st.markdown(
            '<div class="qa-empty">Добавьте первый сайт в разделе «Мои сайты», чтобы начать вести ежеквартальный SEO workspace.</div>',
            unsafe_allow_html=True,
        )
        return

    totals = _status_totals(saved_checks)
    total_cells = len(_all_checks()) * len(sites)
    filled_cells = len(saved_checks)

    metric_cols = st.columns(4)
    metric_cols[0].metric("Сайтов в workspace", len(sites))
    metric_cols[1].metric("Проверок", len(_all_checks()))
    metric_cols[2].metric("Заполнено ячеек", f"{filled_cells}/{total_cells}")
    metric_cols[3].metric("Нужно исправить", totals["needs_fix"])

    date_cols = st.columns([1, 1.3, 1])
    with date_cols[0]:
        bulk_date = st.date_input(
            "Дата проверки",
            value=st.session_state.get("quarterly_bulk_checked_at", date.today()),
            key="quarterly_bulk_checked_at",
            format="DD.MM.YYYY",
        )
    with date_cols[1]:
        bulk_site_id = st.selectbox(
            "Сайт",
            options=[site[0] for site in sites],
            format_func=lambda site_id: _site_name(sites_by_id[site_id]),
            key="quarterly_bulk_site_id",
        )
    with date_cols[2]:
        st.write("")
        st.write("")
        if st.button("Применить дату ко всем проверкам", use_container_width=True):
            _apply_date_to_site(
                user_id=user_id,
                site=sites_by_id[bulk_site_id],
                selected_date=bulk_date,
                saved_checks=saved_checks,
                user_name=user_name,
            )
            st.success("Дата проверки сохранена для всех проверок выбранного сайта.")
            st.rerun()

    selected_sites = st.multiselect(
        "Колонки таблицы",
        options=[site[0] for site in sites],
        default=[site[0] for site in sites],
        format_func=lambda site_id: _site_name(sites_by_id[site_id]),
        help="Можно временно скрыть часть сайтов, сами данные при этом не удаляются.",
    )

    visible_sites = [sites_by_id[site_id] for site_id in selected_sites]
    if not visible_sites:
        st.info("Выберите хотя бы один сайт, чтобы показать таблицу.")
        return

    if st.button("Подготовить Excel", use_container_width=True):
        st.session_state["quarterly_export_data"] = _export_quarterly_excel(
            visible_sites,
            saved_checks,
            cached_get_quarterly_audit_history(user_id=user_id, limit=500),
        )
        st.session_state["quarterly_export_filename"] = (
            f"quarterly_seo_workspace_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        )

    if st.session_state.get("quarterly_export_data"):
        st.download_button(
        "Скачать Excel",
            data=st.session_state["quarterly_export_data"],
            file_name=st.session_state.get("quarterly_export_filename", "quarterly_seo_workspace.xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    st.markdown(
        '<div class="qa-helper">После F5 статусы, комментарии, даты и score-метрики останутся на месте. Excel выгружается по выбранным колонкам сайтов.</div>',
        unsafe_allow_html=True,
    )

    # Legacy tabs were replaced by the radio below so inactive sections do not render.

    quarterly_sections = ["SEO workspace", "History"]
    selected_quarterly_section = st.radio(
        "Раздел quarterly audit",
        quarterly_sections,
        horizontal=True,
        label_visibility="collapsed",
    )

    if selected_quarterly_section == quarterly_sections[0]:
        st.markdown('<div class="qa-workspace-header">', unsafe_allow_html=True)
        header_cols = st.columns([1.15] + [1 for _ in visible_sites], gap="medium")
        with header_cols[0]:
            st.markdown('<div class="qa-workspace-head-cell">SEO проверка</div>', unsafe_allow_html=True)
        for header_col, site in zip(header_cols[1:], visible_sites):
            with header_col:
                st.markdown(
                    f'<div class="qa-workspace-head-cell">{escape(_site_name(site))}</div>',
                    unsafe_allow_html=True,
                )
        st.markdown("</div>", unsafe_allow_html=True)

        _render_workspace_section(
            "Основной блок",
            MAIN_CHECKS,
            visible_sites,
            sites_by_id,
            saved_checks,
            user_id,
            user_name,
        )

        if st.checkbox("Показать ежемесячный блок", value=False):
            _render_workspace_section(
                "Ежемесячный аудит",
                MONTHLY_CHECKS,
                visible_sites,
                sites_by_id,
                saved_checks,
                user_id,
                user_name,
            )

    elif selected_quarterly_section == quarterly_sections[1]:
        st.markdown(
            '<div class="qa-helper">История хранит старые и новые значения статуса и комментария для каждой измененной проверки.</div>',
            unsafe_allow_html=True,
        )
        _render_history(user_id, sites_by_id)
